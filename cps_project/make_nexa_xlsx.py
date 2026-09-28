import sys; sys.path.insert(0,'.')
from nexa_chains import CHAINS, SCAN_V1
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

A="Arial"
HF=PatternFill("solid",fgColor="1F3864"); H=Font(name=A,size=10,bold=True,color="FFFFFF")
BODY=Font(name=A,size=10); BOLD=Font(name=A,size=10,bold=True)
TITLE=Font(name=A,size=14,bold=True,color="1F3864"); NOTE=Font(name=A,size=9,italic=True,color="595959")
MONO=Font(name="Consolas",size=9); STEPF=Font(name=A,size=10)
thin=Side(style="thin",color="BFBFBF"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
SEV={"Medium":PatternFill("solid",fgColor="FFE7CE"),"Low":PatternFill("solid",fgColor="EDEDED"),
     "Informational":PatternFill("solid",fgColor="E2EFF7")}
ENG={"SAST":PatternFill("solid",fgColor="DBE5F1"),"IaC (KICS)":PatternFill("solid",fgColor="E2EFDA")}
STAT={"Confirmed in scan cc434dd7":PatternFill("solid",fgColor="D9EAD3")}
CONF=PatternFill("solid",fgColor="D9EAD3"); PEND=PatternFill("solid",fgColor="FFF2CC")

wb=Workbook()
def hdr(ws,cols,r=4):
    for i,h in enumerate(cols,1):
        c=ws.cell(row=r,column=i,value=h); c.fill=HF; c.font=H; c.border=BOX
        c.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[r].height=30
def W(ws,w):
    for i,x in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=x

# ---------- Sheet 1: Chains ----------
ws=wb.active; ws.title="Chains"
ws["A1"]="Nexa Commerce - vulnerability chain catalogue"; ws["A1"].font=TITLE
ws["A2"]=("Every finding originates in the Nexa Commerce codebase we built - a single, self-contained polyglot "
 "e-commerce application. No third-party or borrowed project is used. Scan cc434dd7 (github.com/cxsmtp/CPS-DEMO "
 "@ 1a27a7d) confirmed 9 of the participants; the rest are pending the v2 re-scan.")
ws["A2"].font=NOTE; ws["A2"].alignment=Alignment(wrap_text=True,vertical="top")
ws.merge_cells("A2:I2"); ws.row_dimensions[2].height=42
hdr(ws,["Chain","Name","Service","Stack","Engines involved","Findings","Severity mix","Chain CPS","Outcome"])
r=5
for c in CHAINS:
    engines=sorted({f[1] for f in c["findings"]})
    from collections import Counter
    mix=Counter(f[3] for f in c["findings"])
    mixs=" ".join(f"{v}x{k[:3] if k!='Informational' else 'Info'}" for k,v in
                  sorted(mix.items(),key=lambda x:{"Medium":0,"Low":1,"Informational":2}[x[0]]))
    vals=[c["id"],c["name"],c["service"],c["stack"]," + ".join(engines),len(c["findings"]),mixs,c["cps"],c["outcome"]]
    for i,v in enumerate(vals,1):
        cell=ws.cell(row=r,column=i,value=v); cell.border=BOX
        cell.font=BOLD if i in (1,8) else BODY
        cell.alignment=Alignment(vertical="top",wrap_text=(i in (2,5,9)))
    ws.cell(row=r,column=8).number_format="0.00"
    ws.cell(row=r,column=8).fill=PatternFill("solid",fgColor="FBD5D5")
    ws.row_dimensions[r].height=42
    r+=1
ws.cell(row=r+1,column=2,value="All ten chains reach the High band. No constituent finding is rated High or Critical.").font=BOLD
W(ws,[8,42,20,26,22,9,16,10,60]); ws.freeze_panes="A5"

# ---------- Sheet 2: Findings & engines (traceability) ----------
ws=wb.create_sheet("Findings & Engines")
ws["A1"]="Findings with detecting engine, severity and traceability"; ws["A1"].font=TITLE
ws["A2"]=("One row per chain participant. 'Scan engine' is the Checkmarx engine that detects it; 'Checkmarx result ID' "
 "locates a confirmed finding in scan cc434dd7 via getFindingDetails.")
ws["A2"].font=NOTE; ws.merge_cells("A2:J2")
hdr(ws,["Chain","Ref","Scan engine","Vulnerability name (as Checkmarx emits it)","Severity","File","Line",
        "Checkmarx result ID","Validation status","CPS-relevant role"])
r=5
ROLE={"F1":"","F2":"","F3":"","F4":"","F5":""}
for c in CHAINS:
    for f in c["findings"]:
        ref,eng,q,sev,fil,ln,rid,stat=f
        vals=[c["id"],ref,eng,q,sev,fil,ln,rid or "-",stat,""]
        for i,v in enumerate(vals,1):
            cell=ws.cell(row=r,column=i,value=v); cell.border=BOX
            cell.font=MONO if i in (4,6,8) else BODY
            cell.alignment=Alignment(vertical="top",wrap_text=(i in (4,6)))
        ws.cell(row=r,column=3).fill=ENG.get(eng,PatternFill())
        ws.cell(row=r,column=5).fill=SEV.get(sev,PatternFill())
        ws.cell(row=r,column=9).fill=CONF if stat.startswith("Confirmed") else PEND
        r+=1
# summary counts
ws.cell(row=r+1,column=3,value="SAST rows:").font=BOLD
ws.cell(row=r+1,column=4,value=f'=COUNTIF(C5:C{r-1},"SAST")').font=BOLD
ws.cell(row=r+2,column=3,value="IaC (KICS) rows:").font=BOLD
ws.cell(row=r+2,column=4,value=f'=COUNTIF(C5:C{r-1},"IaC (KICS)")').font=BOLD
ws.cell(row=r+3,column=3,value="Confirmed in scan:").font=BOLD
ws.cell(row=r+3,column=4,value=f'=COUNTIF(I5:I{r-1},"Confirmed in scan cc434dd7")').font=BOLD
W(ws,[8,6,14,46,13,44,10,32,26,10]); ws.freeze_panes="A5"; ws.auto_filter.ref=f"A4:J{r-1}"

# ---------- Sheet 3: Attack Paths (the centrepiece) ----------
ws=wb.create_sheet("Attack Paths")
ws["A1"]="Step-by-step exploitation - how the gaps in each chain compose"; ws["A1"].font=TITLE
ws["A2"]=("For each chain: the ordered attacker actions, what each step abuses, and which finding(s) enable it. "
 "This is the exploit narrative - how individually low-rated gaps combine into a High-band outcome.")
ws["A2"].font=NOTE; ws.merge_cells("A2:D2"); ws.row_dimensions[2].height=28
r=4
for c in CHAINS:
    # chain banner
    ws.merge_cells(start_row=r,start_column=1,end_row=r,end_column=4)
    b=ws.cell(row=r,column=1,value=f"{c['id']}  {c['name']}   -   {c['stack']}   -   Chain CPS {c['cps']} (High)")
    b.font=Font(name=A,size=11,bold=True,color="FFFFFF"); b.fill=HF
    b.alignment=Alignment(vertical="center"); ws.row_dimensions[r].height=22
    r+=1
    o=ws.cell(row=r,column=1,value="Outcome:"); o.font=BOLD
    ws.merge_cells(start_row=r,start_column=2,end_row=r,end_column=4)
    oc=ws.cell(row=r,column=2,value=c["outcome"]); oc.font=BODY
    oc.alignment=Alignment(wrap_text=True,vertical="top"); r+=1
    for col,h in enumerate(["Step","Action","What it exploits / how","Finding"],1):
        cell=ws.cell(row=r,column=col,value=h); cell.fill=PatternFill("solid",fgColor="D9E2F3")
        cell.font=BOLD; cell.border=BOX
    r+=1
    for n,(title,desc,refs) in enumerate(c["path"],1):
        step="" if title=="Outcome" else str(n)
        ws.cell(row=r,column=1,value=step).font=BODY
        tc=ws.cell(row=r,column=2,value=title); tc.font=BOLD if title!="Outcome" else Font(name=A,size=10,bold=True,italic=True)
        tc.alignment=Alignment(wrap_text=True,vertical="top")
        dc=ws.cell(row=r,column=3,value=desc); dc.font=BODY; dc.alignment=Alignment(wrap_text=True,vertical="top")
        fc=ws.cell(row=r,column=4,value=refs); fc.font=MONO; fc.alignment=Alignment(vertical="top")
        for col in range(1,5): ws.cell(row=r,column=col).border=BOX
        if title=="Outcome":
            for col in range(1,5): ws.cell(row=r,column=col).fill=PatternFill("solid",fgColor="FCE9E9")
        r+=1
    r+=1  # gap between chains
W(ws,[6,30,104,10])
ws.sheet_view.showGridLines=False

# ---------- Sheet 4: Validation ----------
ws=wb.create_sheet("Validation & Method")
ws["A1"]="How these chains are validated"; ws["A1"].font=TITLE
rows=[
 ("Consolidated codebase",""),
 ("Repository","github.com/cxsmtp/CPS-DEMO, subdirectory nexa-commerce/ (commit 1a27a7d scanned as v1; v2 patches pending push)."),
 ("Principle","Every finding in every chain originates in code we wrote for this specimen. No borrowed or third-party project contributes any finding. The specimen is a working polyglot e-commerce application - five runnable services plus Terraform, Kubernetes and Docker."),
 ("",""),
 ("Design constraint",""),
 ("No High or Critical","The application is built so that no chain participant rates High or Critical. Injection is avoided (parameterised or constant SQL, escaped output), the only traversal is the stored-relative variant, every runtime Dockerfile sets USER, and cloud storage is encrypted with public access blocked. Each chain reaches the High band purely by composition of Medium, Low and Informational findings."),
 ("",""),
 ("Scan engines",""),
 ("SAST","37 of 46 findings - application-code weaknesses across PHP, Java, Node, Go and Python."),
 ("IaC (KICS)","9 of 46 findings - Docker Compose, Dockerfile, Terraform and Kubernetes misconfigurations."),
 ("Confirmed in scan cc434dd7","9 findings reproduced at the predicted severity and file in the first scan. The remainder are pending the v2 re-scan, which reshapes the patterns the v1 scan showed were not yet firing."),
 ("",""),
 ("Scoring",""),
 ("Per-finding CPS","(10/4) x weighted sum of five 0-4 dimensions: Prevalence 0.15, Chain Utility 0.30, AI Leverage 0.25, Blast Radius 0.15, Impact Proximity 0.15."),
 ("Chain CPS","max(individual scores) + 0.1 x sum(remaining), capped at 10.0. Bands: 0-2.5 Negligible, 2.6-5.0 Low, 5.1-7.5 Moderate, 7.6-10.0 High."),
 ("",""),
 ("Reproducing",""),
 ("Scan","./scan.sh Nexa-Commerce main  (SAST, SCA, IaC, containers)"),
 ("Verify","PYTHONPATH=/path/to/cps_project python verify_chains.py <export>.json - prints per-chain assembly and flags any High/Critical."),
]
r=3
for a,b in rows:
    ca=ws.cell(row=r,column=1,value=a); ca.font=BOLD if b=="" and a else BODY
    cb=ws.cell(row=r,column=2,value=b); cb.font=BODY; cb.alignment=Alignment(wrap_text=True,vertical="top")
    r+=1
W(ws,[26,110])

wb.save("/mnt/user-data/outputs/Nexa_Commerce_Chains.xlsx")
print("workbook written")
