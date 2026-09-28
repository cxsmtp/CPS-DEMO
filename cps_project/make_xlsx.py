import json
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

D = json.load(open('/tmp/wb.json'))
chains, rows = D["chains"], D["rows"]

ARIAL = "Arial"
HDR_FILL = PatternFill("solid", fgColor="1F3864")
HDR_FONT = Font(name=ARIAL, size=10, bold=True, color="FFFFFF")
BODY = Font(name=ARIAL, size=10)
BOLD = Font(name=ARIAL, size=10, bold=True)
TITLE = Font(name=ARIAL, size=14, bold=True, color="1F3864")
NOTE = Font(name=ARIAL, size=9, italic=True, color="595959")
MONO = Font(name="Consolas", size=9)
thin = Side(style="thin", color="BFBFBF")
BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
SEV_FILL = {"Medium": PatternFill("solid", fgColor="FFE7CE"),
            "Low": PatternFill("solid", fgColor="EDEDED"),
            "Information": PatternFill("solid", fgColor="E2EFF7")}
BAND_FILL = {"High": PatternFill("solid", fgColor="FBD5D5"),
             "Moderate": PatternFill("solid", fgColor="FFF2CC")}

wb = Workbook()

def header(ws, headers, r=1):
    for i, h in enumerate(headers, 1):
        c = ws.cell(row=r, column=i, value=h)
        c.fill, c.font, c.border = HDR_FILL, HDR_FONT, BOX
        c.alignment = Alignment(vertical="center", wrap_text=True)
    ws.row_dimensions[r].height = 30

def widths(ws, w):
    for i, x in enumerate(w, 1):
        ws.column_dimensions[get_column_letter(i)].width = x

# ---------------- Sheet 1: Summary ----------------
ws = wb.active
ws.title = "Summary"
ws["A1"] = "Chain Potential Score - Validated Chain Catalogue"
ws["A1"].font = TITLE
ws["A2"] = ("Every constituent finding was read from a completed Checkmarx One scan in tenant "
            "checkmarx-global-services-internal. Finding IDs are re-checkable via "
            "getFindingDetails(scan_id, finding_id). No chain contains a High or Critical finding.")
ws["A2"].font = NOTE
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:J2"); ws.row_dimensions[2].height = 28

H = ["Chain ID","Chain name","Stack","Findings","Medium","Low","Info","Chain CPS","Band","Source scan"]
header(ws, H, 4)
r = 5
for c in chains:
    m = c["sev_mix"]
    ws.cell(row=r, column=1, value=c["id"]).font = BOLD
    ws.cell(row=r, column=2, value=c["name"]).font = BODY
    ws.cell(row=r, column=3, value=c["stack"]).font = BODY
    ws.cell(row=r, column=4, value=f'=COUNTIF(Traceability!$A:$A,$A{r})')
    for col, key in ((5,"Medium"),(6,"Low"),(7,"Information")):
        ws.cell(row=r, column=col,
                value=f'=COUNTIFS(Traceability!$A:$A,$A{r},Traceability!$F:$F,"{key}")')
    ws.cell(row=r, column=8, value=c["cps"]).number_format = "0.00"
    bc = ws.cell(row=r, column=9, value=c["band"])
    bc.fill = BAND_FILL.get(c["band"], PatternFill())
    ws.cell(row=r, column=10, value=c["scan"]).font = MONO
    for col in range(1, 11):
        ws.cell(row=r, column=col).border = BOX
    for col in (5, 6, 7):
        ws.cell(row=r, column=col).font = BODY
    ws.cell(row=r, column=9).font = BODY
    ws.cell(row=r, column=8).font = BOLD
    r += 1

tot = r
ws.cell(row=tot, column=2, value="TOTAL").font = BOLD
ws.cell(row=tot, column=4, value=f"=SUM(D5:D{r-1})").font = BOLD
for col in (5,6,7):
    L = get_column_letter(col)
    ws.cell(row=tot, column=col, value=f"=SUM({L}5:{L}{r-1})").font = BOLD
ws.cell(row=tot, column=9, value=f'=COUNTIF(I5:I{r-1},"High")&" of "&COUNTA(I5:I{r-1})&" High band"').font = BOLD
for col in range(1, 11):
    ws.cell(row=tot, column=col).border = BOX

n = tot + 2
ws.cell(row=n, column=1, value="Constituent findings rated High or Critical:").font = BOLD
ws.cell(row=n, column=5,
        value=f'=COUNTIFS(Traceability!$F:$F,"High")+COUNTIFS(Traceability!$F:$F,"Critical")').font = BOLD
ws.cell(row=n+1, column=1,
        value=("A finding may serve in more than one chain - that is composition, not duplication. "
               "The Findings sheet shows how many chains use each. Unique findings: see Findings sheet.")
        ).font = NOTE
widths(ws, [11, 52, 22, 10, 9, 8, 8, 11, 11, 40])
ws.freeze_panes = "A5"

# ---------------- Sheet 2: Traceability ----------------
ws = wb.create_sheet("Traceability")
ws["A1"] = "Chain-to-finding traceability matrix"
ws["A1"].font = TITLE
ws["A2"] = ("One row per chain participant. 'Checkmarx finding ID' + 'Scan ID' locate the finding "
            "in the Checkmarx One report; 'Location' locates it in source.")
ws["A2"].font = NOTE
H = ["Chain ID","Chain name","Stack","Chain role","Query name (as Checkmarx emits it)","Severity",
     "Engine","Individual CPS","CPS band","Project","Scan ID","Checkmarx finding ID","Location (file:line)"]
header(ws, H, 4)
r = 5
for x in rows:
    vals = [x["chain"], x["chain_name"], x["stack"], x["role"].replace("_"," "), x["query"],
            x["sev"], x["engine"], x["cps"], x["band"], x["project"], x["scan"], x["fid"], x["loc"]]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = MONO if i in (5, 11, 12, 13) else BODY
        c.border = BOX
        c.alignment = Alignment(vertical="top", wrap_text=(i in (2, 13)))
    ws.cell(row=r, column=6).fill = SEV_FILL.get(x["sev"], PatternFill())
    ws.cell(row=r, column=8).number_format = "0.00"
    r += 1
widths(ws, [10, 42, 20, 14, 46, 12, 10, 12, 11, 34, 40, 34, 60])
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:M{r-1}"

# ---------------- Sheet 3: Findings ----------------
ws = wb.create_sheet("Findings")
ws["A1"] = "Unique findings used across the catalogue"
ws["A1"].font = TITLE
ws["A2"] = "Deduplicated by Checkmarx finding ID. 'Chains using' counts participation across the catalogue."
ws["A2"].font = NOTE
H = ["Checkmarx finding ID","Query name","Severity","Engine","Individual CPS","CPS band",
     "Project","Scan ID","Location (file:line)","Chains using"]
header(ws, H, 4)
seen, r = {}, 5
for x in rows:
    if x["fid"] in seen: continue
    seen[x["fid"]] = True
    vals = [x["fid"], x["query"], x["sev"], x["engine"], x["cps"], x["band"],
            x["project"], x["scan"], x["loc"]]
    for i, v in enumerate(vals, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = MONO if i in (1, 2, 8, 9) else BODY
        c.border = BOX
        c.alignment = Alignment(vertical="top", wrap_text=(i == 9))
    ws.cell(row=r, column=3).fill = SEV_FILL.get(x["sev"], PatternFill())
    ws.cell(row=r, column=5).number_format = "0.00"
    ws.cell(row=r, column=10, value=f'=COUNTIF(Traceability!$L:$L,$A{r})').border = BOX
    r += 1
ws.cell(row=r, column=1, value="Unique findings").font = BOLD
ws.cell(row=r, column=5, value=f"=COUNTA(A5:A{r-1})").font = BOLD
ws.cell(row=r, column=10, value=f"=SUM(J5:J{r-1})").font = BOLD
widths(ws, [34, 46, 12, 10, 12, 11, 34, 40, 60, 12])
ws.freeze_panes = "A5"
ws.auto_filter.ref = f"A4:J{r-1}"

# ---------------- Sheet 4: Scans ----------------
ws = wb.create_sheet("Scans")
ws["A1"] = "Source scan provenance"
ws["A1"].font = TITLE
ws["A2"] = ("Tenant: checkmarx-global-services-internal. All scans completed; findings read "
            "via the Checkmarx MCP on 2026-08-01.")
ws["A2"].font = NOTE
NROWS = 4 + len(rows)
H = ["Project","Scan ID","Stack","Scan totals (C / H / M / L / Info)","Findings used","Chains sourced"]
header(ws, H, 4)
TOTALS = {
 "cx-andy-schmit/dvwa": ("PHP","410 total - 59 C, 104 H, 86 M, 45 L, 116 Info"),
 "cx-carolyn-yates/JavaVulnerableLab": ("Java","860 total - 124 C, 119 H, 178 M, 133 L, 306 Info"),
 "OWASP/NodeGoat Demo": ("JavaScript / Docker","687 total - 84 C, 300 H, 246 M, 55 L, 2 Info"),
 "AISC_openai_agents_python": ("Python / AI agent","48 total - 16 C, 7 H, 1 M, 23 L, 1 Info"),
 "owasp-juice-lab": ("Terraform / AWS / K8s","55 total - 3 C, 12 H, 16 M, 15 L, 9 Info"),
 "cx-jeremy-polansky/authlab-canary": ("Go","57 total - 12 C, 24 H, 17 M, 4 L"),
}
r = 5
for proj, (stack, tot_s) in TOTALS.items():
    scan = next(x["scan"] for x in rows if x["project"] == proj)
    for i, v in enumerate([proj, scan, stack, tot_s], 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = MONO if i == 2 else BODY
        c.border = BOX
    ws.cell(row=r, column=5, value=f'=COUNTIFS(Findings!$G:$G,$A{r})').border = BOX
    ws.cell(row=r, column=6, value=(
        f'=SUMPRODUCT((Traceability!$J$5:$J${NROWS}=$A{r})/'
        f'COUNTIFS(Traceability!$J$5:$J${NROWS},Traceability!$J$5:$J${NROWS},'
        f'Traceability!$A$5:$A${NROWS},Traceability!$A$5:$A${NROWS}))')).border = BOX
    r += 1
ws.cell(row=r+1, column=1, value=(
  "The 'Scan totals' column is the argument: a severity-ordered backlog reaches these chains only "
  "after clearing every Critical and High above them. For the Informational-only chain (CH-106) "
  "that means clearing all 554 Critical/High/Medium/Low findings first.")).font = NOTE
widths(ws, [38, 40, 22, 52, 15, 16])

# ---------------- Sheet 5: Method ----------------
ws = wb.create_sheet("Method")
ws["A1"] = "Method, scoring model and verification"
ws["A1"].font = TITLE
lines = [
 ("", ""),
 ("Scoring model", ""),
 ("Each finding is scored 0-4 on five dimensions, weighted, then normalised to 0-10.", ""),
 ("Dimension", "Weight"),
 ("Prevalence", 0.15), ("Chain Utility", 0.30), ("AI Leverage", 0.25),
 ("Blast Radius", 0.15), ("Impact Proximity", 0.15),
 ("", ""),
 ("Per-finding CPS", "(10/4) x weighted sum of the five dimension scores"),
 ("Chain CPS", "max(individual scores) + 0.1 x sum(remaining scores), capped at 10.0"),
 ("Bands", "0-2.5 Negligible | 2.6-5.0 Low | 5.1-7.5 Moderate | 7.6-10.0 High"),
 ("", ""),
 ("Method", ""),
 ("1. Survey", "Query inventory read live from the tenant across PHP, Java, JavaScript, Go and Python presets, plus KICS and SCA."),
 ("2. Select", "Chains assembled only from findings that already existed, at the severity Checkmarx had already assigned. No prediction step."),
 ("3. Score", "Five-dimension rubric, aggregated by the chain formula above."),
 ("4. Verify", "Each chain re-matched against its own scan's findings. A guard test fails the build if any High or Critical enters a chain. A negative control confirms each fixture assembles only its own chains."),
 ("", ""),
 ("Verification", ""),
 ("Automated tests", "36 covering parser shapes, scorer dispatch, chain assembly, path scoping and the no-High guard."),
 ("Re-check any row", "getFindingDetails(scan_id, finding_id) using the Traceability sheet columns K and L."),
 ("", ""),
 ("Notes on composition", ""),
 ("Shared findings", "A finding may participate in several chains. Composition, not duplication - the Findings sheet counts participation."),
 ("Moderate-band chains", "Three chains land in the Moderate band rather than High. They are retained deliberately: the framework grades chains, it does not assert that every chain is critical."),
 ("Severity is preset-dependent", "The same query can carry different severities in different language presets - e.g. Use of Insufficiently Random Values is Medium in PHP and Low in JavaScript. Severity is bound to (query, language)."),
 ("Query names vary", "Casing and wording vary across presets and even within one scan. Matching requires normalisation."),
]
r = 3
for a, b in lines:
    ca = ws.cell(row=r, column=1, value=a)
    ca.font = BOLD if b == "" and a else BODY
    cb = ws.cell(row=r, column=2, value=b)
    cb.font = BODY
    cb.alignment = Alignment(wrap_text=True, vertical="top")
    if isinstance(b, float): cb.number_format = "0.00"
    r += 1
widths(ws, [30, 110])

# ---------------- Sheet 6: Specimen reproduction ----------------
ws = wb.create_sheet("Specimen Reproduction")
ws["A1"] = "Nexa Commerce - single-codebase reproduction"
ws["A1"].font = TITLE
ws["A2"] = ("The chains above were observed across six separate applications. Nexa Commerce is a "
            "purpose-built polyglot e-commerce specimen intended to reproduce them in one repository. "
            "Status below is from scan cc434dd7 (repo github.com/cxsmtp/CPS-DEMO, commit 1a27a7d, v1).")
ws["A2"].font = NOTE
ws["A2"].alignment = Alignment(wrap_text=True, vertical="top")
ws.merge_cells("A2:E2"); ws.row_dimensions[2].height = 40
H = ["Chain","Query reproduced","Severity observed","Location in specimen","Status"]
header(ws, H, 4)
REPRO = [
 ("CH-101","Insecure_Value_of_the_SameSite_Cookie_Attribute","Medium","storefront/public/cart.php:25","Reproduced"),
 ("CH-101","Cookie_Overly_Broad_Path","Low","storefront/public/cart.php:25","Reproduced"),
 ("CH-102","Information_Exposure_Through_an_Error_Message","Low","catalog-service/.../CatalogServlet.java:34","Reproduced"),
 ("CH-103","External_Control_of_System_or_Config_Setting","Medium","catalog-service/.../ConfigServlet.java:44","Reproduced"),
 ("CH-103","Creation_of_Temp_File_With_Insecure_Permissions","Low","catalog-service/.../ReportBuilder.java:59","Reproduced (query name differs from source scan)"),
 ("CH-104","Open_Redirect","Medium","web-gateway/src/server.js:86","Reproduced"),
 ("CH-104","Log_Forging","Low","web-gateway/src/server.js:69, 85","Reproduced"),
 ("CH-105","Use_Of_Hardcoded_Password","Medium","ops/seed/seed.js:15","Reproduced"),
 ("CH-110","Client_Weak_Cryptographic_Hash","Medium","auth-service/public/js/checkout.js:54","Reproduced"),
]
r = 5
for row_ in REPRO:
    for i, v in enumerate(row_, 1):
        c = ws.cell(row=r, column=i, value=v)
        c.font = MONO if i in (2, 4) else BODY
        c.border = BOX
        c.alignment = Alignment(vertical="top", wrap_text=(i in (4, 5)))
    ws.cell(row=r, column=3).fill = SEV_FILL.get(row_[2], PatternFill())
    r += 1
r += 1
for t, b in [
  ("Reproduced in v1", "9 of 48 constituent findings, each at the predicted severity in the predicted file."),
  ("Defects found in v1", "Two, both in the specimen and both fixed in v2: a Critical Stored_XSS from streaming a template file to the HTTP response, and a High Plain_Text_Transport_Layer_in_Server from a non-TLS Go listener."),
  ("Cause of misses (1)", "Python used str.format() where the reference framework uses f-strings; Checkmarx taint tracking follows the f-string interpolation."),
  ("Cause of misses (2)", "Java read input via request.getAttribute(), which is not a taint source. Corrected to request.getParameter()."),
  ("Cause of misses (3)", "PHP patterns sat in lib/ includes; in this tenant the PHP preset reports at the entry-point page, not inside the include."),
  ("Status", "v2 patches written and packaged; awaiting push and re-scan.")]:
    ws.cell(row=r, column=1, value=t).font = BOLD
    c = ws.cell(row=r, column=2, value=b); c.font = BODY
    c.alignment = Alignment(wrap_text=True, vertical="top")
    r += 1
widths(ws, [22, 52, 18, 46, 44])
ws.freeze_panes = "A5"

wb.save("/mnt/user-data/outputs/CPS_Chain_Catalogue.xlsx")
print("workbook written")
