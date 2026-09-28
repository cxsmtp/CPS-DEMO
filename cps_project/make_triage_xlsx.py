from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SCAN = "cc434dd7-3e6b-4df0-8ccf-b43ac3ec7678"
PROJ = "d182f9ec-bb1e-40fb-9b79-12dbb96f7f1c"

ARIAL="Arial"
HDR_FILL=PatternFill("solid", fgColor="1F3864")
HDR=Font(name=ARIAL,size=10,bold=True,color="FFFFFF")
BODY=Font(name=ARIAL,size=10); BOLD=Font(name=ARIAL,size=10,bold=True)
TITLE=Font(name=ARIAL,size=14,bold=True,color="1F3864")
NOTE=Font(name=ARIAL,size=9,italic=True,color="595959")
MONO=Font(name="Consolas",size=9)
thin=Side(style="thin",color="BFBFBF"); BOX=Border(left=thin,right=thin,top=thin,bottom=thin)
NE=PatternFill("solid",fgColor="D9EAD3")     # proposed not exploitable
KEEP=PatternFill("solid",fgColor="FCE5CD")   # keep / true positive
DUP=PatternFill("solid",fgColor="FFF2CC")    # duplicate

# (result_id, similarity_id, query, sev, file, line, verdict, rationale, action)
TRIAGE = [
 ("v4e5wHWcmrMX0TQpVHlhdJNGeIM=","1357008053","Stored_XSS","CRITICAL",
  "/nexa-commerce/web-gateway/src/server.js","76 -> 82","False positive",
  "Checkmarx models the flow readFileSync (line 76) -> body -> res.end (line 82) as a stored-XSS: data read from a "
  "persistent store and returned unsanitised. The store here is src/views/product.html, a static template baked into "
  "the image at build time by 'COPY src ./src'. No code path in the service writes to the views directory; the "
  "container runs as an unprivileged user with no request data reaching the filesystem. Stored XSS requires an "
  "attacker-writable store, and none exists. The only interpolated values come from the in-module CATALOGUE constant "
  "and are passed through escapeHtml(). Exploitation would first require arbitrary file write to the image, which is "
  "a different and absent vulnerability.",
  "NOT_EXPLOITABLE"),
 ("la6h3sIYYd9KF6i8D5VbpcMqIgo=","848664396","Unchecked_Input_for_Loop_Condition","MEDIUM",
  "/nexa-commerce/verify_chains.py","47 -> 56","False positive",
  "The tainted source is json.loads() of the file named by the operator's own --catalog argument, and the sink is a "
  "for-loop over that catalogue. verify_chains.py is offline research tooling invoked manually from a terminal; it is "
  "not imported by any service, has no network listener, is not copied into any container image, and is excluded from "
  "the build. The 'attacker' able to set the loop bound is the operator running the script against their own file. No "
  "trust boundary is crossed.",
  "NOT_EXPLOITABLE"),
 ("nTUWpXsPM0wyT439V1/0/Ghp1n0=","508863670","Unchecked_Input_for_Loop_Condition","MEDIUM",
  "/nexa-commerce/verify_chains.py","47 -> 54","False positive",
  "Same source and same file as the finding above, reported at a second loop sink (line 54, iteration over "
  "catalog['chains']). Identical reasoning: operator-supplied local JSON in a non-deployed CLI utility, no trust "
  "boundary, no network surface.",
  "NOT_EXPLOITABLE"),
 ("XVOwzk3eDmete1xqbjQapOtNL+o=","630543548","Unchecked_Input_for_Loop_Condition","MEDIUM",
  "/nexa-commerce/verify_chains.py","47 -> 90","False positive",
  "Same source and file, third loop sink (line 90, iteration over the assembled 'rows' list). Identical reasoning. "
  "The three findings differ only in which loop the analyser terminated the path at.",
  "NOT_EXPLOITABLE"),
]

KEEPERS = [
 ("I/7OE/sSBnmbSNHMJiMcWr6utmU=","External_Control_of_System_or_Config_Setting","MEDIUM",
  "/nexa-commerce/catalog-service/.../ConfigServlet.java","44","True positive - intended",
  "Deliberate chain participant CH-103 F1. A request parameter is written into a JVM-wide system property. Real and "
  "intended; keep as-is."),
 ("TQPeKdcuM1U7VfYPBF/sSGwS0D4=","Environment_Variable_Injection","MEDIUM",
  "/nexa-commerce/catalog-service/.../ConfigServlet.java","44","Duplicate view",
  "Same sink and same line as External_Control_of_System_or_Config_Setting above - two queries reporting one "
  "System.setProperty call. Genuinely reachable, so NOT triaged away, but it should be counted once, not twice, when "
  "tallying distinct issues. Note the query name is imprecise: setProperty sets a JVM system property, not an OS "
  "environment variable."),
 ("JC8IH8Ezh474NyadGwbv+sAebck=","Open_Redirect","MEDIUM",
  "/nexa-commerce/web-gateway/src/server.js","86","True positive - intended",
  "Deliberate chain participant CH-104 F1. return_to is written into the Location header with no allow-list."),
 ("tHYT8FVsSDsnAHGCwgpL/SEZZBI=","Use_Of_Hardcoded_Password","MEDIUM",
  "/nexa-commerce/ops/seed/seed.js","15","True positive - intended",
  "Deliberate chain participant CH-105 F1. Operational database password compiled into the seed image."),
 ("LP9LVOzghsVbI1vYCJxrCHR23n4=","Insecure_Value_of_the_SameSite_Cookie_Attribute","MEDIUM",
  "/nexa-commerce/storefront/public/cart.php","25","True positive - intended",
  "Deliberate chain participant CH-101 F3."),
 ("LrfnsakFDfvvkYiwdzwMVCVT8Fc=","Cookie_Overly_Broad_Path","LOW",
  "/nexa-commerce/storefront/public/cart.php","25","True positive - intended",
  "Deliberate chain participant CH-101 F4."),
 ("n6Ek0rV+jtRmdWCufmCtVqcK5dw=","Log_Forging","LOW",
  "/nexa-commerce/web-gateway/src/server.js","69","True positive - intended",
  "Deliberate chain participant CH-104 F5. Raw query value written to the application log."),
 ("qQRyGisPw858N/fsTSL5KdN9pJg=","Log_Forging","LOW",
  "/nexa-commerce/web-gateway/src/server.js","85","True positive - intended",
  "Second instance of CH-104 F5 on the sign-in path."),
 ("UHUxxVQtJ43G2rAPTE/+et6h+v0=","Information_Exposure_Through_an_Error_Message","LOW",
  "/nexa-commerce/catalog-service/.../CatalogServlet.java","34","True positive - intended",
  "Deliberate chain participant CH-102 F4. printStackTrace writes to the response body."),
 ("sE+dyQtToO+tPv4lx2b4meetWLU=","Creation_of_Temp_File_With_Insecure_Permissions","LOW",
  "/nexa-commerce/catalog-service/.../ReportBuilder.java","59","True positive - intended",
  "Deliberate chain participant CH-103 F4. createTempFile in the shared temp directory."),
 ("ZFzh5id20NAGkCS/srpi0QR5yrM=","Trust_Boundary_Violation_in_Session_Variables","LOW",
  "/nexa-commerce/assistant-service/src/server.py","31","True positive - not in a chain",
  "The session identifier is taken from the query string and used as the session key. Real weakness, correctly "
  "reported; not part of a declared chain. Left TO_VERIFY."),
 ("VYnQ284PzKZEqutU2Mt+uuqUn0U=","Unchecked_Input_for_Loop_Condition","MEDIUM",
  "/nexa-commerce/assistant-service/src/agents/mcp/util.py","26","True positive - keep",
  "Unlike the verify_chains.py instances, this loop iterates the tool list from a decoded MCP envelope, which is "
  "attacker-influenceable if the MCP server or the transport is compromised. A real trust boundary is crossed. NOT "
  "triaged as a false positive."),
 ("Y3U2Vepx9Ay8MwF0+vHIYcklJO0=","Generation of Error Message Containing Sensitive Information","MEDIUM",
  "/nexa-commerce/catalog-service/.../admin/dashboard.jsp","19","True positive - intended",
  "The JSP declares no errorPage and web.xml declares no error-page, so an exception from InventoryDao renders the "
  "container trace. This is the mechanism behind CH-106 F3."),
 ("PW6CydZD4nVV/LTJOK/nky5CrWo=","Missing_HSTS_Header","MEDIUM",
  "/nexa-commerce/catalog-service/.../admin/dashboard.jsp","1","True positive",
  "No Strict-Transport-Security is written by the catalog service. Real; keep."),
]

wb=Workbook()

def header(ws,H,r=4):
    for i,h in enumerate(H,1):
        c=ws.cell(row=r,column=i,value=h); c.fill=HDR_FILL; c.font=HDR; c.border=BOX
        c.alignment=Alignment(vertical="center",wrap_text=True)
    ws.row_dimensions[r].height=32
def widths(ws,w):
    for i,x in enumerate(w,1): ws.column_dimensions[get_column_letter(i)].width=x

# Sheet 1 - proposed triage
ws=wb.active; ws.title="Proposed Triage"
ws["A1"]="False-positive triage - proposed NOT_EXPLOITABLE"; ws["A1"].font=TITLE
ws["A2"]=(f"Scan {SCAN} | Project CPS-DEMO {PROJ} | repo github.com/cxsmtp/CPS-DEMO @ 1a27a7d. "
          "Every row below is a SAST finding; SAST and SCA are the only engines Checkmarx One accepts triage for.")
ws["A2"].font=NOTE; ws["A2"].alignment=Alignment(wrap_text=True,vertical="top")
ws.merge_cells("A2:I2"); ws.row_dimensions[2].height=28
H=["#","Checkmarx result ID","Similarity ID","Query","Severity","File","Line (src -> sink)","Verdict","Triage reasoning","Proposed state"]
header(ws,H)
r=5
for i,(rid,sim,q,sev,f,ln,verdict,why,act) in enumerate(TRIAGE,1):
    for col,v in enumerate([i,rid,sim,q,sev,f,ln,verdict,why,act],1):
        c=ws.cell(row=r,column=col,value=v); c.border=BOX
        c.font=MONO if col in (2,3,6) else BODY
        c.alignment=Alignment(vertical="top",wrap_text=(col in (6,9)))
    ws.cell(row=r,column=10).fill=NE; ws.cell(row=r,column=10).font=BOLD
    ws.cell(row=r,column=8).fill=NE
    ws.row_dimensions[r].height=96
    r+=1
ws.cell(row=r+1,column=1,value="Proposed NOT_EXPLOITABLE:").font=BOLD
ws.cell(row=r+1,column=4,value=f'=COUNTIF(J5:J{r-1},"NOT_EXPLOITABLE")').font=BOLD
widths(ws,[5,34,14,44,10,44,18,16,86,18])
ws.freeze_panes="A5"

# Sheet 2 - keep / true positives
ws=wb.create_sheet("Confirmed & Keep")
ws["A1"]="Findings reviewed and retained"; ws["A1"].font=TITLE
ws["A2"]=("Reviewed in the same pass and deliberately NOT triaged away. Most are intended chain participants; "
          "the remainder are genuine weaknesses outside any declared chain.")
ws["A2"].font=NOTE; ws.merge_cells("A2:G2")
H=["#","Checkmarx result ID","Query","Severity","File","Line","Verdict","Reasoning"]
header(ws,H)
r=5
for i,(rid,q,sev,f,ln,verdict,why) in enumerate(KEEPERS,1):
    for col,v in enumerate([i,rid,q,sev,f,ln,verdict,why],1):
        c=ws.cell(row=r,column=col,value=v); c.border=BOX
        c.font=MONO if col in (2,5) else BODY
        c.alignment=Alignment(vertical="top",wrap_text=(col in (5,8)))
    ws.cell(row=r,column=7).fill = DUP if "Duplicate" in verdict else KEEP
    ws.row_dimensions[r].height=60
    r+=1
widths(ws,[5,34,48,10,46,8,26,90])
ws.freeze_panes="A5"

# Sheet 3 - scope notes
ws=wb.create_sheet("Scope Notes")
ws["A1"]="Scope, engine coverage and what was deliberately not triaged"; ws["A1"].font=TITLE
notes=[
 ("Scan composition",""),
 ("Total findings","7,903 across SAST, SCA, IaC (KICS) and Containers."),
 ("Containers","7,640 findings - base-image CVEs from the tomcat, node, python, golang, alpine and php images. "
  "Not triaged: Checkmarx One accepts triage only for SAST and SCA, and these are inherited from upstream images "
  "rather than written code. They are out of scope for the specimen's no-High claim."),
 ("cps_project/ directory","The repository also contains the CPS research labs, whose findings are deliberate "
  "research artefacts (Path_Traversal, vulnerable Flask/Werkzeug pins, privilege-escalation manifests). They are "
  "genuine true positives by design and were NOT marked not-exploitable. To evaluate the specimen alone, scan the "
  "nexa-commerce/ subdirectory, or exclude cps_project/** at the project level."),
 ("",""),
 ("Triage engine support",""),
 ("Supported","sast, sca"),
 ("Not supported","kics, containers, secret-detection - the Checkmarx One triage API does not yet accept state "
  "changes for these engines, so no IaC or container finding could be triaged even where a case exists."),
 ("",""),
 ("Method",""),
 ("Evidence","Every verdict was reached by reading the finding's full data flow via getFindingDetails, not from the "
  "title alone. Source and sink lines are recorded in the Proposed Triage sheet."),
 ("Standard applied","A finding is marked not exploitable only where the data flow requires a precondition that does "
  "not exist in this codebase - an attacker-writable store, or a trust boundary that is never crossed. Findings that "
  "are merely low-impact, intended, or inconvenient were left alone."),
 ("Reversibility","Triage state is per-finding and reversible in the Checkmarx One UI. Comments are recorded against "
  "each finding and readable later via getTriageHistory."),
]
r=3
for a,b in notes:
    ca=ws.cell(row=r,column=1,value=a); ca.font=BOLD if b=="" and a else BODY
    cb=ws.cell(row=r,column=2,value=b); cb.font=BODY
    cb.alignment=Alignment(wrap_text=True,vertical="top")
    if b: ws.row_dimensions[r].height=max(15, 15*(len(b)//95+1))
    r+=1
widths(ws,[24,112])

wb.save("/mnt/user-data/outputs/CPS_Triage_FalsePositives.xlsx")
print("triage workbook written")
